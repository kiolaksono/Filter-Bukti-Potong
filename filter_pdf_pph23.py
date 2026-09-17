"""
Filter PDF Bukti Potong berdasarkan Nomor Bukti Potong di Excel
==================================================================

Cara pakai:
1. Taruh file ini di satu folder yang sama dengan:
   - Semua file PDF Bukti Potong yang mau difilter.
   - Satu (atau lebih) file Excel (.xlsx) yang berisi kolom
     "NOMOR BUKTI POTONG" (nama kolom boleh beda huruf besar/kecil
     atau ada spasi, program tetap mengenalinya).
2. Jalankan (double klik, atau lewat terminal: python filter_bukpot.py).
3. Program akan:
   - Membaca semua Nomor Bukti Potong dari file Excel.
   - Membaca Nomor Bukti Potong dari SETIAP file PDF di folder
     (dibaca langsung dari isi PDF-nya, bukan dari nama file).
   - Kalau Nomor Bukti Potong PDF tsb COCOK dengan salah satu yang ada
     di Excel, file PDF dipindahkan ke folder "Filtered".
   - PDF yang tidak cocok akan tetap di folder asal (tidak diubah).
4. Setelah selesai, ringkasan hasil (berapa dipindah, berapa nomor di
   Excel yang tidak ketemu PDF-nya, dst) ditampilkan di layar dan
   dicatat ke file "Filter_Log.txt" di folder yang sama.

Requirement: pip install pdfplumber openpyxl
"""

import re
import sys
import shutil
from pathlib import Path

import pdfplumber
import openpyxl

# ----------------------------------------------------------------------
# Konfigurasi
# ----------------------------------------------------------------------

FILTERED_FOLDER_NAME = "Filtered"
LOG_FILENAME = "Filter_Log.txt"
TARGET_COLUMN_NAMES = {"nomor bukti potong", "no bukti potong", "nomor bukpot"}

# File Excel yang dibuat oleh program lain (rekap) diabaikan supaya
# tidak ikut dibaca sebagai sumber daftar filter.
EXCLUDED_EXCEL_PREFIXES = ("rekap_", "filter_")


# ----------------------------------------------------------------------
# Baca daftar Nomor Bukti Potong dari semua file Excel di folder
# ----------------------------------------------------------------------

def normalize_code(s):
    if s is None:
        return None
    return str(s).strip().upper()


def read_target_codes(excel_files):
    """Return dict {kode_ternormalisasi: kode_asli} dari semua file Excel."""
    codes = {}
    sources = []
    for xlsx_path in excel_files:
        try:
            wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        except Exception as e:
            print(f"  [PERINGATAN] Gagal membuka {xlsx_path.name}: {e}")
            continue

        found_in_file = 0
        for ws in wb.worksheets:
            # cari baris header (baris pertama yang ada isinya) dan cari
            # kolom yang namanya cocok dengan "Nomor Bukti Potong"
            header_row = None
            col_idx = None
            for row in ws.iter_rows(min_row=1, max_row=min(5, ws.max_row)):
                for cell in row:
                    if cell.value and normalize_code(cell.value).lower() in TARGET_COLUMN_NAMES:
                        header_row = cell.row
                        col_idx = cell.column
                        break
                if col_idx:
                    break
            if not col_idx:
                continue

            for row in ws.iter_rows(min_row=header_row + 1, min_col=col_idx, max_col=col_idx):
                val = row[0].value
                if val is None or str(val).strip() == "":
                    continue
                norm = normalize_code(val)
                codes[norm] = str(val).strip()
                found_in_file += 1

        if found_in_file:
            sources.append((xlsx_path.name, found_in_file))

    return codes, sources


# ----------------------------------------------------------------------
# Baca Nomor Bukti Potong dari isi PDF
# ----------------------------------------------------------------------

def get_nomor_bukti_potong(pdf_path: Path):
    with pdfplumber.open(pdf_path) as pdf:
        text = pdf.pages[0].extract_text() or ""

    # Baris header BPPU: "<NOMOR> <MASA PAJAK> <SIFAT> <STATUS>"
    m = re.search(
        r"^(\S+)\s+(\d{2}-\d{4})\s+(TIDAK\s+FINAL|FINAL)\s+(.+)$",
        text, re.MULTILINE,
    )
    if m:
        return m.group(1).strip()
    return None


# ----------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------

def main():
    if getattr(sys, "frozen", False):
        folder = Path(sys.executable).resolve().parent
    else:
        folder = Path(__file__).resolve().parent

    log_lines = []

    def log(msg=""):
        print(msg)
        log_lines.append(msg)

    pdf_files = sorted(folder.glob("*.pdf"))
    excel_files = [
        f for f in sorted(folder.glob("*.xlsx"))
        if not f.name.lower().startswith(EXCLUDED_EXCEL_PREFIXES)
    ]

    if not excel_files:
        log("Tidak ada file Excel (.xlsx) ditemukan di folder ini.")
        input("Tekan ENTER untuk keluar...")
        return

    if not pdf_files:
        log(f"Tidak ada file PDF ditemukan di folder ini ({folder}).")
        input("Tekan ENTER untuk keluar...")
        return

    log(f"File Excel ditemukan: {', '.join(f.name for f in excel_files)}")
    target_codes, sources = read_target_codes(excel_files)

    if not target_codes:
        log('Tidak ditemukan kolom "NOMOR BUKTI POTONG" yang bisa dibaca di file Excel.')
        input("Tekan ENTER untuk keluar...")
        return

    for name, count in sources:
        log(f"  - {name}: {count} nomor bukti potong")
    log(f"Total nomor bukti potong unik di Excel: {len(target_codes)}")
    log(f"Total file PDF di folder: {len(pdf_files)}")
    log("")

    filtered_dir = folder / FILTERED_FOLDER_NAME
    matched_codes = set()
    moved = 0
    skipped = 0
    unreadable = []

    for pdf_path in pdf_files:
        try:
            nomor = get_nomor_bukti_potong(pdf_path)
        except Exception as e:
            unreadable.append((pdf_path.name, str(e)))
            log(f"  [GAGAL BACA] {pdf_path.name} -> {e}")
            continue

        if nomor is None:
            unreadable.append((pdf_path.name, "Nomor Bukti Potong tidak terbaca dari PDF"))
            log(f"  [GAGAL BACA] {pdf_path.name} -> Nomor Bukti Potong tidak terbaca dari PDF")
            continue

        norm = normalize_code(nomor)
        if norm in target_codes:
            filtered_dir.mkdir(exist_ok=True)
            dest = filtered_dir / pdf_path.name
            shutil.move(str(pdf_path), str(dest))
            matched_codes.add(norm)
            moved += 1
            log(f"  [PINDAH] {pdf_path.name}  (Nomor Bukti Potong: {nomor})")
        else:
            skipped += 1

    # nomor di Excel yang tidak ketemu PDF-nya sama sekali
    not_found = [orig for norm, orig in target_codes.items() if norm not in matched_codes]

    log("")
    log("=" * 60)
    log(f"Selesai. {moved} file dipindahkan ke folder '{FILTERED_FOLDER_NAME}'.")
    log(f"{skipped} file PDF tidak cocok dengan Excel (dibiarkan di folder asal).")
    if unreadable:
        log(f"{len(unreadable)} file PDF gagal dibaca nomornya (dibiarkan di folder asal).")
    if not_found:
        log(f"\n{len(not_found)} Nomor Bukti Potong di Excel TIDAK ditemukan PDF-nya:")
        for code in not_found:
            log(f"  - {code}")

    (folder / LOG_FILENAME).write_text("\n".join(log_lines), encoding="utf-8")
    log(f"\nLog lengkap disimpan di: {folder / LOG_FILENAME}")
    input("\nTekan ENTER untuk keluar...")


if __name__ == "__main__":
    main()
